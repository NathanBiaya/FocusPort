import os
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.text import PP_ALIGN

def create_presentation():
    prs = Presentation()
    
    # This finds your Desktop path automatically
    desktop = os.path.join(os.path.join(os.environ['USERPROFILE']), 'Desktop') 
    file_path = os.path.join(desktop, 'FocusPort_Presentation.pptx')

    content = [
        {"title": "NATHAN: SOFTWARE DEVELOPER & PHOTOGRAPHER", "notes": "Hello, my name is Nathan. I am a CS student at CSCC..."},
        {"title": "VISION: MINIMALIST PHOTOGRAPHY SHOWCASE", "notes": "FocusPort is a distraction-free portfolio..."},
        {"title": "EVOLUTION: LAYOUT, DYNAMICS, DEPLOYMENT", "notes": "Phase 1: Structure. Phase 2: Logic. Phase 3: Cloud..."},
        {"title": "ADAPTATION: SOLVING DATA PERSISTENCE", "notes": "We overcame ephemeral storage failures by migrating to Postgres..."},
        {"title": "FUTURE: SCALING WITH SPECIALIZED TEAMS", "notes": "I would assemble a team with UI and Security specialists..."}
    ]

    for item in content:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = item["title"]
        slide.notes_slide.notes_text_frame.text = item["notes"]

    prs.save(file_path)
    print(f"DONE! Look on your Desktop for: FocusPort_Presentation.pptx")

if __name__ == "__main__":
    create_presentation()